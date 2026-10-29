"""The import screens show the format of their file: the columns, what each holds, and a blank format file."""
import io

import openpyxl


def test_import_formats(client, admin_headers):
    csv = client.get("/api/v1/meta/import-format/members", headers=admin_headers)
    assert csv.status_code == 200, csv.text
    cols = {c["name"]: c for g in csv.json()["groups"] for c in g["columns"]}
    assert cols["membership_number"]["required"] and cols["member_name"]["required"] and not cols["email"]["required"]
    heads = client.get("/api/v1/meta/import-template/members", headers=admin_headers).text.strip().split(",")
    assert set(heads) == set(cols), "the format and the template list the same columns"
    assert client.get("/api/v1/meta/import-format/nope", headers=admin_headers).status_code == 404

    reg = client.get("/api/v1/members/legacy-import/format", headers=admin_headers)
    assert reg.status_code == 200, reg.text
    rcols = {c["name"]: c for g in reg.json()["groups"] for c in g["columns"]}
    assert rcols["CusNo"]["required"] and rcols["CusName"]["required"]
    assert "No means no magazine" in rcols["CusMagazine"]["meaning"] and "CusStatus" in rcols and "CusExpDt" in rcols
    tpl = client.get("/api/v1/members/legacy-import/template", headers=admin_headers)
    assert tpl.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(tpl.content))
    assert wb.sheetnames == ["MstMember", "Format"] and wb["MstMember"].max_row == 1
    assert {"CusNo", "CusName", "CusMagazine", "CusReceiptNo6"} <= {c.value for c in wb["MstMember"][1]}
